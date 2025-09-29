from docx import Document
from docx.shared import Pt

# LibreOffice ondersteunt .docx goed, dus dit is compatibel.
doc = Document()

# Functie om titel toe te voegen
def add_heading(text, level=1):
    doc.add_heading(text, level=level)

# Functie om normale tekst toe te voegen
def add_paragraph(text):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(11)

# Titel van het document
add_heading("Boekstructuur – Brigitte Schrijvers", level=0)

# Deel I
add_heading("DEEL I – Overleven", level=1)
add_paragraph("Dit deel gaat over de donkere periode, het kind dat overleefde, de pijnlijke waarheden, het ontstaan van binnenwerelden.")

add_heading("Hoofdstukken:", level=2)
add_paragraph("1. De kunst van overleven – Over leven met trauma, zonder houvast, zonder stem.")
add_paragraph("2. De stilte tussen de klappen – Hoe zwijgen soms luider is dan schreeuwen.")
add_paragraph("3. Fantasievriendjes & binnenwerelden – Creatieve overleving, een veilige plek bouwen.")
add_paragraph("4. Wanneer God je rug aanraakt – De mystieke ervaringen, en het wonder dat je bleef leven.")

# Deel II
add_heading("DEEL II – Ontwaken", level=1)
add_paragraph("De eerste stappen naar bewustwording, genezing, zelfliefde en zachtheid.")

add_heading("Hoofdstukken:", level=2)
add_paragraph("5. Slaap als zelfliefde – Hoe rust een daad van erkenning werd.")
add_paragraph("6. Pijn in de voeten, maar licht in de ziel – Leren lopen ondanks de pijn.")
add_paragraph("7. De val, de vonk, de wederopstanding – Het dieptepunt dat een keerpunt werd.")
add_paragraph("8. Waarom ik toch blijf geven – Helpen als missie, ondanks alles.")

# Deel III
add_heading("DEEL III – Leven", level=1)
add_paragraph("De lessen, het licht, het relativeren en leren leven – ook met humor.")

add_heading("Hoofdstukken:", level=2)
add_paragraph("9. De kracht van een glimlach die niemand zag – Onzichtbare daden van kracht.")
add_paragraph("10. Wakker worden in een slapende wereld – Wat je ziet als je zelf ontwaakt.")
add_paragraph("11. Leren relativeren als levenskunst – Humor, zelfspot en de kracht van omdenken.")
add_paragraph("12. Slot: De les die overbleef – De eindconclusie als zachte samenvatting.")

# Kleurenpalet suggesties
add_heading("Kleurenpalet per deel (symbolisch):", level=1)
add_paragraph("• Deel I – Overleven: Grijstinten, donkerblauw, paars (rouw, innerlijke diepte)")
add_paragraph("• Deel II – Ontwaken: Zachtgeel, lichtblauw, roze (herstel, hoop, zachtheid)")
add_paragraph("• Deel III – Leven: Regenboog, groen, goud (groei, licht, volledigheid)")

# Studentennotities voorbeeld
add_heading("Voorbeeld ‘Studentennotities’ – Relativerende humor:", level=1)
add_paragraph("Studentennotitie 1: Therapie? Dat is praten tegen iemand die jou begrijpt, terwijl jij jezelf niet volgt – met koffie erbij.")
add_paragraph("Studentennotitie 2: Mijn jeugd was als een IKEA-kast zonder handleiding. Alles zat erin, maar niks paste.")

# Bestand opslaan
file_path = "/mnt/data/Boekstructuur_Brigitte_Schrijvers.docx"
doc.save(file_path)

file_path
